# -*- coding: utf-8 -*-
"""
test_doc_ops.py — 文档读写 / 清洗 / 解析类修复的回归测试

覆盖：
  · 1框 回写：只改差异字符，制表符 / 软回车 / 上标 / 超链接原位保留
  · 半角→全角逐处判断、括号成对、跨 run 的「1.」序号、全角数字不算中文
  · 修正按 occurrence 只改一处；occurrence 计数与定位一致
  · 附图标记段识别、标记名含数字 / 多编号共用名称
  · 「摘要：正文」同段、权项主题名称保护、词库跨词误报、更新说明数组

运行：
    python tests/test_doc_ops.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def _para(inner_xml: str):
    doc = Document()
    p = doc.add_paragraph()
    for el in parse_xml(f'<w:p {nsdecls("w")}>{inner_xml}</w:p>'):
        p._p.append(el)
    return p


class _Shell:
    def __init__(self, t):
        self.text = t


def test_writeback_keeps_structure_and_formatting():
    from core.paragraph_edit import display_text, set_paragraph_text
    p = _para('<w:r><w:t>1.</w:t><w:tab/><w:t>一种装置</w:t></w:r>')
    assert display_text(p) == "1.\t一种装置"
    assert set_paragraph_text(p, "1.\t一种新装置")
    assert p.text == "1.\t一种新装置", p.text
    assert not set_paragraph_text(p, "1.\t一种新装置"), "未改动的段不应被判为有变化"

    p = _para('<w:r><w:t>第一行</w:t><w:br/><w:t>第二行</w:t></w:r>')
    assert display_text(p) == "第一行↵第二行", "软回车显示为 ↵，不拆行"
    set_paragraph_text(p, "第一行改↵第二行")
    assert p.text == "第一行改\n第二行", repr(p.text)

    p = _para('<w:r><w:t>面积为10m</w:t></w:r>'
              '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>2</w:t></w:r>'
              '<w:r><w:t>，壳体</w:t></w:r>')
    set_paragraph_text(p, "面积为20m2，外壳体")
    sup_run = p.runs[1]
    assert sup_run.text == "2" and sup_run.font.superscript, "上标 run 应原样保留"

    p = _para('<w:r><w:t>前</w:t></w:r><w:hyperlink><w:r><w:t>链接</w:t></w:r></w:hyperlink>'
              '<w:r><w:t>后</w:t></w:r>')
    assert display_text(p) == p.text == "前链接后"
    set_paragraph_text(p, "前链接文字后")
    assert p.text == "前链接文字后"
    print("[OK] 回写只改差异字符，制表符 / 软回车 / 上标 / 超链接保留")


def test_halfwidth_punct_per_position():
    from core.cleaner import unify_halfwidth_punct
    cases = {
        "采用公式f(x,y)计算,结果为1:2": "采用公式f(x,y)计算，结果为1:2",
        "所述壳体(1)的端部": "所述壳体（1）的端部",
        "(1)壳体": "（1）壳体",
        "如下:壳体;底座": "如下：壳体；底座",
    }
    for src, want in cases.items():
        p = _para(f'<w:r><w:t>{src}</w:t></w:r>')
        unify_halfwidth_punct([p])
        assert p.text == want, f"{src!r} → {p.text!r}，期望 {want!r}"

    p = _para('<w:r><w:t>1</w:t></w:r><w:r><w:t>.一种装置.所述</w:t></w:r>')
    unify_halfwidth_punct([p])
    assert p.text == "1.一种装置。所述", f"跨 run 的序号不应被改：{p.text!r}"
    p = _para('<w:r><w:t>１.一种装置</w:t></w:r>')
    unify_halfwidth_punct([p])
    assert p.text == "１.一种装置", f"全角数字序号不应被改：{p.text!r}"
    print("[OK] 半角→全角逐处判断，括号成对，序号不误改")


def test_corrections_apply_only_their_occurrence():
    from core.cleaner import apply_typo_corrections, occurrence_at, nth_occurrence
    text = "权力要求书和权力要求"
    pos = text.rfind("权力要求")
    occ = occurrence_at(text, "权力要求", pos)
    assert occ == 2 and nth_occurrence(text, "权力要求", occ) == pos, "occurrence 计数须与定位一致"

    p = _para('<w:r><w:t>所述的的壳体与所述的的底座</w:t></w:r>')
    n = apply_typo_corrections([p], [
        {"para_idx": 0, "wrong": "的的", "confirmed_fix": "的", "occurrence": 2},
    ])
    assert n == 1 and p.text == "所述的的壳体与所述的底座", p.text
    p = _para('<w:r><w:t>的的和的的</w:t></w:r>')
    apply_typo_corrections([p], [{"para_idx": 0, "wrong": "的的", "confirmed_fix": "的"}])
    assert p.text == "的和的", "不带 occurrence 时仍按全部出现替换"
    print("[OK] 修正按 occurrence 只改一处")


def test_duplicate_words_report_each_occurrence():
    from core.cleaner import check_duplicate_words
    res = check_duplicate_words([_Shell("所述所述壳体与所述所述底座")])
    assert [r["occurrence"] for r in res if r["wrong"] == "所述所述"] == [1, 2], res
    print("[OK] 重复字词每处各报一条并带 occurrence")


def test_wordbank_no_cross_word_hits():
    from core.cleaner import check_typos_wordbank
    texts = ["所述上限为10mm", "所述圆形装置", "开设至少一个孔", "与轴成45度",
             "作用与反作用", "连接件部分", "气流成环状", "北京技术有限公司"]
    hits = check_typos_wordbank([_Shell(t) for t in texts])
    assert not hits, [(h["context"], h["wrong"]) for h in hits]
    assert check_typos_wordbank([_Shell("所述权力要求书")]), "正常错字仍应检出"
    print("[OK] 词库不再跨词误命中")


def _make_docx(lines):
    tmpdir = tempfile.mkdtemp()
    path = os.path.join(tmpdir, "t.docx")
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    return path


def test_mark_paragraph_and_inline_abstract():
    from core.doc_parser import parse_document
    from core.spec_check import check_abstract_length
    path = _make_docx([
        "权利要求书",
        "1.一种夹具，其特征在于，包括齿圈。",
        "技术领域",
        "本发明涉及夹具。",
        "附图说明",
        "图1为本发明的结构示意图；",
        "图2、图3分别为左视图和俯视图。",
        "附图标记：",
        "1-齿圈，2-夹指；",
        "3-转盘。",
        "具体实施方式",
        "所述齿圈固定。",
        "摘要：本发明公开了一种夹具。" + "很长的摘要" * 70,
    ])
    d = parse_document(path)
    assert [p.text for p in d["mark_paras"]] == ["附图标记：", "1-齿圈，2-夹指；", "3-转盘。"], \
        [p.text for p in d["mark_paras"]]
    sec = d["sections"]["说明书摘要"]
    assert sec.end_idx > sec.start_idx and "本发明公开了" in d["paragraphs"][sec.start_idx].text
    res = check_abstract_length(d["paragraphs"], d["sections"])
    body = "本发明公开了一种夹具。" + "很长的摘要" * 70
    assert res and f"约 {len(body)} 字" in res[0]["message"], res
    print("[OK] 附图标记段跳过「图2、图3」；摘要标题与正文同段可识别、前缀不计字数")


def test_mark_names():
    from core.mark_extractor import extract_marks_from_text as f
    assert f("10-壳体，11、12-侧板，13-底座") == {10: "壳体", 11: "侧板", 12: "侧板", 13: "底座"}
    assert f("2-第3连杆，4-壳体") == {2: "第3连杆", 4: "壳体"}
    assert f("1-齿圈2-夹指") == {1: "齿圈", 2: "夹指"}
    assert f("附图标记：1、底座；2、第一支架；30、电机") == {1: "底座", 2: "第一支架", 30: "电机"}
    print("[OK] 标记名含数字 / 多编号共用名称")


def test_claims_subject_protection():
    from core.annotator import annotate_section
    from core.doc_parser import DocSection
    paras = [_para('<w:r><w:t>1.一种齿圈装置，包括齿圈和夹指，其特征在于，所述齿圈转动。</w:t></w:r>')]
    annotate_section(paras, DocSection("权利要求书", 0, 1), {"齿圈": "齿圈（1）", "夹指": "夹指（2）"},
                     skip_already_annotated=False, mode="claims")
    assert paras[0].text == "1.一种齿圈装置，包括齿圈（1）和夹指（2），其特征在于，所述齿圈（1）转动。", paras[0].text
    print("[OK] 只保护主题名称，前序部分的部件照常标注")


def test_update_notes_list():
    from infra.updater import UpdateInfo
    info = UpdateInfo.from_json({"version": "4.4", "notes": ["feat: a", "", "- b"]})
    assert info.notes == "feat: a\n\n- b", repr(info.notes)
    print("[OK] 更新说明为数组时按行拼接")


def test_content_area_edit_safety():
    """1框：回写后显示与内存一致；有未确认编辑的页不被重排覆盖；锁定期间不可确认。"""
    from PyQt6.QtWidgets import QApplication
    from ui.main_window import MainWindow
    app = QApplication.instance() or QApplication(sys.argv)
    path = _make_docx(["权利要求书", "1.一种夹具，其特征在于，包括齿圈。", "技术领域", "本发明涉及夹具。"])
    win = MainWindow()
    win._load_document(path)
    ca = win.content_area
    ed = ca._edits[0]
    lines = ed.toPlainText().split("\n")
    lines[0] = lines[0].replace("齿圈", "齿圈和夹指")
    ed.setPlainText("\n".join(lines))
    assert ca.has_pending_edits()

    ca._rich_tabs.add(0)                  # 模拟含图片页的缩放重排
    ca._relayout_rich_tabs()
    assert "夹指" in ed.toPlainText(), "未确认的编辑不应被重排覆盖"
    ca._rich_tabs.discard(0)

    ca.set_locked(True)
    assert ed.isReadOnly() and not ca._confirm_btn.isEnabled()
    ca.set_locked(False)
    assert not ed.isReadOnly() and ca._confirm_btn.isEnabled()

    assert win._commit_content_edits()
    claim = next(p for p in win.doc_data["paragraphs"] if p.text.startswith("1."))
    assert "齿圈和夹指" in claim.text, claim.text
    assert win._has_unsaved_changes()
    win._unsaved = False
    win.close()
    print("[OK] 1框 编辑：不被重排覆盖、锁定、提交前置")


def test_confirm_marks_rewrites_all_mark_paragraphs():
    from PyQt6.QtWidgets import QApplication
    from ui.main_window import MainWindow
    app = QApplication.instance() or QApplication(sys.argv)
    path = _make_docx([
        "附图说明", "图1为结构示意图；", "附图标记：1-齿圈，", "2-夹指。",
        "具体实施方式", "所述齿圈固定。",
    ])
    win = MainWindow()
    win._load_document(path)
    win.marks_edit.setPlainText("1-齿圈，2-夹指，3-转盘")
    win._on_confirm_marks()
    texts = [p.text for p in win.doc_data["mark_paras"]]
    assert texts == ["附图标记：1-齿圈，2-夹指，3-转盘", ""], texts
    assert "3-转盘" in win.content_area._edits[1].toPlainText(), "确认后 1框 应刷新"
    win._unsaved = False
    win.close()
    print("[OK] 确认标记：合并写入首段并保留前缀，其余标记段清空，1框 刷新")


if __name__ == "__main__":
    test_writeback_keeps_structure_and_formatting()
    test_halfwidth_punct_per_position()
    test_corrections_apply_only_their_occurrence()
    test_duplicate_words_report_each_occurrence()
    test_wordbank_no_cross_word_hits()
    test_mark_paragraph_and_inline_abstract()
    test_mark_names()
    test_claims_subject_protection()
    test_update_notes_list()
    test_content_area_edit_safety()
    test_confirm_marks_rewrites_all_mark_paragraphs()
    print("\nAll doc-ops tests passed.")
