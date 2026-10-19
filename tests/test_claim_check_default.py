# -*- coding: utf-8 -*-
"""main-new 默认模式漏报回归：两个动态选项始终关闭。"""
import hashlib
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from core.claim_check import run_all_checks


ROOT = Path(__file__).resolve().parents[1]
SIZES = (2, 3, 4, 5, 6, 7, 30)
TERMS = ("齿圈", "第一支架", "上盖", "连接件", "弹性件", "对接头", "接地端子", "从动齿轮")


def check(lines, n=2, **kwargs):
    return run_all_checks(
        [SimpleNamespace(text=line) for line in lines], 0, len(lines), n=n,
        check_vague=False, use_dynamic_truncate=False, use_dynamic_fallback=False,
        **kwargs,
    )


class DefaultClaimTests(unittest.TestCase):
    def test_missing_terms_are_checked_at_every_size(self):
        for term in TERMS:
            for n in SIZES:
                for suffix in ("的外壁设有凹槽", "（2），固定于底座"):
                    with self.subTest(term=term, n=n, suffix=suffix):
                        results = check([f"1.一种装置，包括底座；所述{term}{suffix}。"], n)
                        self.assertEqual(len(results), 1, results)
                        self.assertEqual(results[0]["kind"], "antecedent")
                        self.assertIn(results[0]["term"], term)

    def test_defined_terms_are_not_reported(self):
        for term in TERMS:
            for n in SIZES:
                with self.subTest(term=term, n=n):
                    self.assertEqual(check([
                        f"1.一种装置，包括{term}；所述{term}的外壁设有凹槽。",
                    ], n), [])

    def test_default_mode_does_not_silently_fall_back_to_shorter_name(self):
        for defined, missing in (("齿圈", "齿圈盖"), ("电机", "电机安装座"),
                                 ("第一侧", "第一侧盖")):
            for n in (3, 4, 5, 6, 7, 30):
                with self.subTest(defined=defined, missing=missing, n=n):
                    results = check([
                        f"1.一种装置，包括{defined}；所述{missing}的外壁设有凹槽。",
                    ], n)
                    self.assertEqual(len(results), 1, results)
                    self.assertEqual(results[0]["kind"], "antecedent")
                    self.assertTrue(results[0]["term"].startswith(defined), results)
                    self.assertGreater(len(results[0]["term"]), len(defined))
                    self.assertEqual(check([
                        f"1.一种装置，包括{defined}和{missing}；所述{missing}的外壁设有凹槽。",
                    ], n), [])

    def test_spaces_and_citation_do_not_hide_references(self):
        for prefix in ("所述", "所述的", "所述 ", "所述\u3000", "所述 的 "):
            with self.subTest(prefix=prefix):
                results = check([
                    "1.一种装置，包括底座。",
                    f"2.根据权利要求1所述，{prefix}齿圈的外壁设有凹槽。",
                ], 5)
                self.assertEqual(len(results), 1, results)
                self.assertEqual(results[0]["term"], "齿圈")
                self.assertEqual(results[0]["para_idx"], 1)
                self.assertEqual(results[0]["anchor"], prefix + "齿圈")

    def test_repeated_missing_term_is_reported_once(self):
        # main-new 的约定是只报首处，与保留的 main 测试期望报两次不同。
        results = check([
            "1.一种装置，包括底座；所述齿圈的外壁设有凹槽；",
            "所述齿圈的内壁固定于底座。",
            "2.根据权利要求1所述的装置，所述齿圈的端部设有凹槽。",
        ], 5)
        self.assertEqual(len(results), 1, results)
        self.assertEqual(results[0]["claim_no"], 1)
        self.assertEqual(results[0]["term"], "齿圈")

    def test_default_ignore_and_dedup_do_not_hide_longer_names(self):
        lines = ["1.一种装置，包括电机；所述电机安装座的外壁设有凹槽；"
                 "所述电机安装板的外壁设有凹槽；所述电机安装座与底座连接。"]
        results = check(lines, 5, ignore_set={"电机"})
        self.assertEqual([r["term"] for r in results], ["电机安装座", "电机安装板"])
        results = check(lines, 5, ignore_set={"电机安装座"})
        self.assertEqual([r["term"] for r in results], ["电机安装板"])

    def test_no_parsed_claims_is_not_success(self):
        for lines in ([], [""], ["一种装置，所述齿圈的外壁设有凹槽。"]):
            with self.subTest(lines=lines):
                results = check(lines)
                self.assertEqual(len(results), 1, results)
                self.assertEqual(results[0]["kind"], "parse")

    def test_preserved_word_sample_and_inserted_errors(self):
        samples = sorted(ROOT.glob("*初稿*.docx"))
        if not samples:
            self.skipTest("本地未提供保留的 Word 样例")
        from core.doc_parser import parse_document
        sample = samples[0]
        digest = hashlib.sha256(sample.read_bytes()).digest()
        data = parse_document(str(sample))
        section = data["sections"]["权利要求书"]
        lines = [p.text for p in data["paragraphs"][section.start_idx:section.end_idx]]
        for n in SIZES:
            with self.subTest(n=n, scenario="original"):
                baseline = check(lines, n)
                if n == 2:
                    self.assertEqual(baseline, [])
            for term in ("对接头", "接地端子", "第一支架", "弹性件"):
                with self.subTest(n=n, term=term):
                    changed = list(lines)
                    changed[1] = f"所述{term}的外壁设有凹槽；" + changed[1]
                    results = check(changed, n)
                    self.assertTrue(any(r["kind"] == "antecedent" and r["para_idx"] == 1
                                        and r["term"] in term for r in results), results)
                    changed[0] += f"还包括{term}；"
                    self.assertEqual([r["message"] for r in check(changed, n)],
                                     [r["message"] for r in baseline])
        self.assertEqual(hashlib.sha256(sample.read_bytes()).digest(), digest)


class DefaultClaimUITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PyQt6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        from ui.main_window import MainWindow
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)
        settings = Mock()
        settings.get_theme.return_value = "light"
        settings.get_geometry.return_value = None
        settings.get_bool.side_effect = lambda key, default=False: default
        for patcher in (
            patch("config.config_manager.AppSettings", return_value=settings),
            patch("config.config_manager.get_config_dir", return_value=self.tmpdir.name),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        self.win = MainWindow()
        self.win._show_toast = Mock()
        self.win._confirm_discard_changes = Mock(return_value=True)
        self.addCleanup(self.win.close)
        self.win.claim_dyn_trunc_cb.setChecked(False)
        self.win.claim_dyn_fb_cb.setChecked(False)

    def test_unparsed_input_warns_and_valid_definition_clears_results(self):
        from docx import Document
        from PyQt6.QtWidgets import QMessageBox
        path = Path(self.tmpdir.name) / "claims.docx"
        doc = Document()
        doc.add_paragraph("权利要求书")
        doc.add_paragraph("1.一种装置，包括底座；所述对接头的外壁设有凹槽。")
        doc.save(path)
        with patch.object(QMessageBox, "critical") as critical:
            self.win._load_document(str(path))
            self.assertTrue(self.win._claim_loaded)
            editor = self.win.content_area._edits[0]
            self.win._claim_n_buttons.button(5).click()
            self.win.claim_check_btn.click()
            self.assertEqual(self.win._claim_results[0]["term"], "对接头")

            editor.setPlainText("一种装置，所述对接头的外壁设有凹槽。")
            self.win.claim_check_btn.click()
            self.assertEqual(self.win.claim_result_table.item(0, 0).text(), "识别异常")
            self.assertIn("检查未完成", self.win.claim_status_label.text())
            self.assertEqual(self.win._show_toast.call_args.args[1], "warning")

            editor.setPlainText("1.一种装置，包括对接头；所述对接头的外壁设有凹槽。")
            self.win.claim_check_btn.click()
            self.assertEqual(self.win._claim_results, [])
            self.assertEqual(self.win.claim_result_table.rowCount(), 0)
            self.assertEqual(editor.extraSelections(), [])
            self.win._show_toast.assert_called_with("未发现问题", "success")
            critical.assert_not_called()

    def test_preserved_sample_edit_check_table_and_location(self):
        samples = sorted(ROOT.glob("*初稿*.docx"))
        if not samples:
            self.skipTest("本地未提供保留的 Word 样例")
        from PyQt6.QtWidgets import QMessageBox
        digest = hashlib.sha256(samples[0].read_bytes()).digest()
        with patch.object(QMessageBox, "critical") as critical:
            self.win._load_document(str(samples[0]))
            self.assertTrue(self.win._claim_loaded)
            editor = self.win.content_area._edits[0]
            lines = editor.toPlainText().split("\n")
            lines[1] = "所述对接头（99）的外壁设有凹槽；" + lines[1]
            editor.setPlainText("\n".join(lines))
            for n in SIZES:
                with self.subTest(n=n):
                    if n <= 7:
                        self.win._claim_n_buttons.button(n).click()
                    else:
                        self.win.claim_n_custom.setValue(n)
                    self.win.claim_check_btn.click()
                    self.assertFalse(self.win.claim_dyn_trunc_cb.isChecked())
                    self.assertFalse(self.win.claim_dyn_fb_cb.isChecked())
                    results = self.win._claim_results
                    row = next(i for i, r in enumerate(results)
                               if r["kind"] == "antecedent" and r["term"] in "对接头")
                    issue = results[row]
                    self.assertEqual(issue["para_idx"], self.win._claim_start_idx + 1)
                    self.assertEqual(self.win.claim_result_table.item(row, 0).text(), "引用基础")
                    self.assertIn("缺少引用基础", self.win.claim_result_table.item(row, 2).text())
                    self.assertEqual(self.win._show_toast.call_args.args[1], "warning")
                    self.win._on_claim_cell_clicked(row, 2)
                    self.assertTrue(self.win.content_area._active_issue)
                    self.assertEqual(editor.textCursor().blockNumber(), 1)
                    self.assertIn(issue["anchor"], [s.cursor.selectedText() for s in editor.extraSelections()])
            critical.assert_not_called()
        self.assertIn("所述对接头", self.win.doc_data["paragraphs"][self.win._claim_start_idx + 1].text)
        self.assertEqual(hashlib.sha256(samples[0].read_bytes()).digest(), digest)

    def test_turning_off_dynamic_options_restores_full_name_check(self):
        from docx import Document
        path = Path(self.tmpdir.name) / "long_name.docx"
        doc = Document()
        doc.add_paragraph("权利要求书")
        doc.add_paragraph("1.一种装置，包括电机；所述电机安装座的外壁设有凹槽。")
        doc.save(path)
        self.win._load_document(str(path))
        self.win._claim_n_buttons.button(5).click()
        for trunc, fallback in ((False, False), (False, True), (True, False),
                                (True, True), (False, False)):
            with self.subTest(trunc=trunc, fallback=fallback):
                self.win.claim_dyn_trunc_cb.setChecked(trunc)
                self.win.claim_dyn_fb_cb.setChecked(fallback)
                self.win.claim_check_btn.click()
                results = self.win._claim_results
                if trunc or fallback:
                    self.assertEqual(results, [])
                else:
                    self.assertEqual([r["term"] for r in results], ["电机安装座"])
                    self.assertEqual(self.win.claim_result_table.rowCount(), 1)
                    self.assertIn("所述电机安装座", self.win.claim_result_table.item(0, 2).text())
        editor = self.win.content_area._edits[0]
        editor.setPlainText("1.一种装置，包括电机和电机安装座；所述电机安装座的外壁设有凹槽。")
        self.win.claim_check_btn.click()
        self.assertEqual(self.win._claim_results, [])
        self.assertEqual(self.win.claim_result_table.rowCount(), 0)


if __name__ == "__main__":
    unittest.main()
