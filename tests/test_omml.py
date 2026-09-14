# -*- coding: utf-8 -*-
"""
test_omml.py — OMML（Office Math）公式预览回归测试

覆盖：
  · ui/render/omml.py 各类结构都能排出非空位图，且版面关系合理
    （分式比单字高、带角标比裸底宽、不认识的结构降级成文字而不是空白）
  · media.iter_content 把 OMML 产出成图片而不是【公式】占位
  · has_renderable_object 对"只有公式、没有图片"的段也返回 True
    —— 否则纯文本路径会让公式在 1框 里彻底消失
  · has_readonly_object 把含 OMML 的段判为只读，挡住整段文本回写污染 w:t
  · _recolor_lineart 只动灰阶线稿（公式预览），彩色附图原样保留

运行：
    python tests/test_omml.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from lxml import etree

_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _omath(body: str):
    """用一段 OMML 片段造出一个 m:oMath 元素。"""
    return etree.fromstring(f'<m:oMath xmlns:m="{_NS}">{body}</m:oMath>')


def _r(text: str) -> str:
    return f"<m:r><m:t>{text}</m:t></m:r>"


def _e(body: str) -> str:
    return f"<m:e>{body}</m:e>"


_CASES = {
    "文本":     _r("x"),
    "分式":     f"<m:f><m:num>{_r('a')}</m:num><m:den>{_r('b')}</m:den></m:f>",
    "下标":     f"<m:sSub>{_e(_r('q'))}<m:sub>{_r('1')}</m:sub></m:sSub>",
    "上标":     f"<m:sSup>{_e(_r('x'))}<m:sup>{_r('2')}</m:sup></m:sSup>",
    "上下标":   f"<m:sSubSup>{_e(_r('I'))}<m:sub>{_r('x')}</m:sub>"
                f"<m:sup>{_r('2')}</m:sup></m:sSubSup>",
    "根式":     f"<m:rad><m:deg/>{_e(_r('2'))}</m:rad>",
    "括号":     f'<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr>'
                f"{_e(_r('σ'))}</m:d>",
    "大算符":   f'<m:nary><m:naryPr><m:chr m:val="∑"/></m:naryPr>'
                f"<m:sub>{_r('i')}</m:sub><m:sup>{_r('n')}</m:sup>{_e(_r('A'))}</m:nary>",
    "矩阵":     f"<m:m><m:mr>{_e(_r('a'))}{_e(_r('b'))}</m:mr>"
                f"<m:mr>{_e(_r('c'))}{_e(_r('d'))}</m:mr></m:m>",
    "函数":     f"<m:func><m:fName>{_r('sin')}</m:fName>{_e(_r('x'))}</m:func>",
    "上划线":   f'<m:bar><m:barPr><m:pos m:val="top"/></m:barPr>{_e(_r("x"))}</m:bar>',
    "多行":     f"<m:eqArr>{_e(_r('a=1'))}{_e(_r('b=2'))}</m:eqArr>",
}


_QAPP = None


def _app():
    """拿到（必要时创建）QApplication，并用模块级变量按住它。

    只 `return QApplication(...)` 而调用方不接收的话，引用计数会立刻归零、
    QApplication 被析构，之后任何 QImage / QFont 调用都会硬崩（0xC0000409）。
    """
    global _QAPP
    from PyQt6.QtWidgets import QApplication
    _QAPP = QApplication.instance() or QApplication(sys.argv)
    return _QAPP


def test_all_constructs_render():
    """每种 OMML 结构都要排出非空位图，且尺寸合理。"""
    _app()
    from ui.render.omml import omml_to_qimage

    sizes = {}
    for name, body in _CASES.items():
        img = omml_to_qimage(_omath(body), color="#202020", px=16)
        assert img is not None and not img.isNull(), f"{name} 渲染失败"
        assert img.width() > 2 and img.height() > 2, f"{name} 尺寸异常 {img.size()}"
        sizes[name] = (img.width(), img.height())

    # 版面关系：分式比单个字符高；带角标比裸底宽；矩阵比单字既宽又高
    assert sizes["分式"][1] > sizes["文本"][1], sizes
    assert sizes["下标"][0] > sizes["文本"][0], sizes
    assert sizes["矩阵"][0] > sizes["文本"][0] and sizes["矩阵"][1] > sizes["文本"][1], sizes
    # ∑ 带上下限要比裸 ∑ 高
    bare = omml_to_qimage(_omath(_r("∑")), color="#202020", px=16)
    assert sizes["大算符"][1] > bare.height(), (sizes["大算符"], bare.height())
    print(f"[OK] {len(_CASES)} 种 OMML 结构全部排版成功，版面关系正确")


def test_unknown_element_falls_back_to_text():
    """不认识的结构降级成文字，绝不返回空白。"""
    _app()
    from ui.render.omml import omml_to_qimage, omml_text

    weird = _omath(f"<m:zzUnknown>{_r('待定')}</m:zzUnknown>")
    img = omml_to_qimage(weird, color="#202020", px=16)
    assert img is not None and img.width() > 4, "未知结构应降级成文字而不是空白"
    assert omml_text(weird) == "待定"
    # 完全没有文字的空公式：允许返回 None，但不能抛异常
    omml_to_qimage(_omath("<m:zzUnknown/>"), color="#202020", px=16)
    print("[OK] 未知 OMML 结构降级为文字，空公式不抛异常")


def _docx_with_omml(path):
    """造一个含 OMML 公式、但不含任何图片的 docx。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph("说明书")
    doc.add_paragraph("技术领域")
    p = doc.add_paragraph("总荷载为：")
    p._p.append(_omath(
        f"<m:sSub>{_e(_r('q'))}<m:sub>{_r('1')}</m:sub></m:sSub>"
    ))
    doc.add_paragraph("具体实施方式")
    doc.add_paragraph("本实施例中的钢模板包括底板。")
    doc.save(path)


def test_omml_paragraph_renders_as_image_not_placeholder():
    """只有公式、没有图片的段：必须走富文本并产出图片，而不是【公式】占位。"""
    import tempfile
    _app()
    from core.doc_parser import parse_document, _has_image, has_readonly_object
    from ui.render.media import iter_content, has_renderable_object

    tmpdir = tempfile.mkdtemp()
    path = os.path.join(tmpdir, "omml.docx")
    _docx_with_omml(path)

    d = parse_document(path)
    doc = d["document"]
    cache = {}
    target = next(p for p in d["paragraphs"] if p.text.startswith("总荷载为"))

    assert has_renderable_object(target, doc, cache), \
        "含 OMML 的段必须走富文本路径，否则公式会整个消失"
    kinds = [k for k, _ in iter_content(target, doc, cache, "#2c3e50", 16)]
    assert "image" in kinds, f"OMML 应产出图片，实际 {kinds}"
    assert "placeholder" not in kinds, f"不应再退化成【公式】占位，实际 {kinds}"

    # 回写守卫：OMML 不是图片，_has_image 认不出来；has_readonly_object 必须认得
    assert not _has_image(target), "前提：这段确实没有图片/OLE 对象"
    assert has_readonly_object(target), \
        "含 OMML 的段必须判为只读，否则整段文本回写会把占位文字写进 w:t"

    try:
        os.remove(path)
        os.rmdir(tmpdir)
    except OSError:
        pass
    print("[OK] 纯 OMML 段落：走富文本 + 产出图片 + 判为只读")


def test_sample_document_has_no_formula_placeholder():
    """仓库里的真实样例：全文不应再残留任何【公式】占位。"""
    _app()
    from core.doc_parser import parse_document
    from ui.render.media import iter_content

    samples = [f for f in os.listdir(".") if f.endswith(".docx") and "初稿" in f]
    if not samples:
        print("[SKIP] 当前目录没有样例 docx")
        return
    d = parse_document(samples[0])
    doc = d["document"]
    cache = {}
    n_img = n_ph = 0
    for p in d["paragraphs"]:
        for kind, _ in iter_content(p, doc, cache, "#2c3e50", 16):
            if kind == "image":
                n_img += 1
            elif kind == "placeholder":
                n_ph += 1
    assert n_img > 0, "样例里应当解析出图片 / 公式"
    assert n_ph == 0, f"样例里仍有 {n_ph} 处【公式】/【图】占位"
    print(f"[OK] 样例 {samples[0][:20]}…：{n_img} 处图片/公式全部可预览，0 占位")


def test_recolor_only_touches_grayscale_lineart():
    """白底黑线的公式预览转主题色；彩色图必须原样保留。"""
    _app()
    from PyQt6.QtGui import QImage, QColor
    from ui.render.media import _recolor_lineart

    # 灰阶线稿：白底 + 一条黑线
    art = QImage(20, 10, QImage.Format.Format_RGB32)
    art.fill(QColor("#ffffff"))
    for x in range(20):
        art.setPixelColor(x, 5, QColor("#000000"))
    out = _recolor_lineart(art, "#e6e6e6")
    assert out is not None, "灰阶线稿应被重新着色"
    assert out.pixelColor(0, 0).alpha() == 0, "白底应变成全透明"
    ink = out.pixelColor(3, 5)
    assert ink.alpha() > 200, "黑线应保持不透明"
    assert ink.red() > 200 and ink.green() > 200, f"线条应换成主题色，实际 {ink.name()}"

    # 彩色图：不碰
    color_img = QImage(20, 10, QImage.Format.Format_RGB32)
    color_img.fill(QColor("#ffffff"))
    color_img.setPixelColor(5, 5, QColor("#ff3300"))
    assert _recolor_lineart(color_img, "#e6e6e6") is None, "彩色图不应被重新着色"
    print("[OK] 线稿重新着色只作用于灰阶公式预览，彩色附图不受影响")


def test_theme_switch_rerenders_formulas():
    """切主题后公式取色跟着变（颜色是烘焙在像素里的，必须重排）。"""
    import tempfile
    _app()
    from core.doc_parser import parse_document
    from ui.content_area import ContentArea

    tmpdir = tempfile.mkdtemp()
    path = os.path.join(tmpdir, "omml2.docx")
    _docx_with_omml(path)

    area = ContentArea()
    area.load(parse_document(path))
    assert area._rich_tabs, "含公式的标签页应走富文本"
    assert area._math_color == ContentArea.MATH_COLOR_LIGHT

    before = len([k for k in area._img_cache if isinstance(k, tuple)])
    area.set_math_color(ContentArea.MATH_COLOR_DARK)
    assert area._math_color == ContentArea.MATH_COLOR_DARK
    assert before > 0, "浅色主题下应已缓存过公式位图"
    # 旧配色的缓存被丢弃、按新配色重建
    assert all(k[-2] == ContentArea.MATH_COLOR_DARK
               for k in area._img_cache if isinstance(k, tuple) and k[0] == "omml"), \
        "缓存里不应残留旧配色的公式位图"

    area.deleteLater()
    try:
        os.remove(path)
        os.rmdir(tmpdir)
    except OSError:
        pass
    print("[OK] 切换主题后公式按新配色重排，旧配色缓存被丢弃")


if __name__ == "__main__":
    test_all_constructs_render()
    test_unknown_element_falls_back_to_text()
    test_omml_paragraph_renders_as_image_not_placeholder()
    test_sample_document_has_no_formula_placeholder()
    test_recolor_only_touches_grayscale_lineart()
    test_theme_switch_rerenders_formulas()
    print("\nAll OMML formula tests passed.")
